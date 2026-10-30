"""Build a Word report: original image next to its door/window detection result.

Usage: python make_detection_report.py
"""
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image, ImageOps
from pillow_heif import register_heif_opener

register_heif_opener()

ROOT = Path(__file__).resolve().parents[2]
REPORTS_DIR = ROOT / "reports"
IMAGE_DIR = ROOT / "01_Data" / "Example of Inventory" / "Images"
RESULT_DIR = ROOT / "reports" / "output_of_SAM_doors_windows"
REPORT_PATH = REPORTS_DIR / "detection_report.docx"
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp", ".heic"}
IMAGE_WIDTH_CM = 8.7

METHOD = (
    "Method: each photo was processed in two steps. First, Grounding DINO "
    "(open-vocabulary detector) located doors and windows from the text prompt "
    "\"door. window.\". Then each detected box was passed to SAM (Segment Anything), "
    "which produced a precise mask. In the result images, coloured overlays mark the "
    "segmented masks and red rectangles show the detected boxes with their label and "
    "confidence score. Images with no overlay had no detection above the thresholds "
    "(box 0.35, text 0.25)."
)


def image_stream(path: Path) -> io.BytesIO:
    img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    img.thumbnail((1200, 1200))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    buf.seek(0)
    return buf


def add_picture(cell, path: Path, caption: str):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(image_stream(path), width=Cm(IMAGE_WIDTH_CM))
    cap = cell.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True


def main():
    images = sorted(
        (p for p in IMAGE_DIR.iterdir() if p.suffix.lower() in IMAGE_SUFFIXES),
        key=lambda p: p.name.lower(),
    )

    doc = Document()
    section = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))
    doc.styles["Normal"].font.size = Pt(9)
    doc.add_heading("Door and window detection results", level=0)
    doc.add_paragraph(METHOD)

    for path in images:
        doc.add_heading(path.name, level=2)
        table = doc.add_table(rows=1, cols=2)
        left, right = table.rows[0].cells
        add_picture(left, path, "Original")

        result = RESULT_DIR / f"{path.stem}.png"
        if result.exists():
            add_picture(right, result, "Detection result (doors and windows)")
        else:
            right.paragraphs[0].add_run("No result image for this photo.")

    doc.save(REPORT_PATH)
    print(f"Saved {REPORT_PATH}")


if __name__ == "__main__":
    main()
