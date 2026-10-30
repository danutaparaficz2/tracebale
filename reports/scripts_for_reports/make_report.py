"""Build a Word report: each image on the left, its vLLM analysis on the right.

Usage: python make_report.py
"""
import io
import json
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from PIL import Image, ImageOps
from pillow_heif import register_heif_opener

register_heif_opener()

ROOT = Path(__file__).resolve().parents[2]
REPORTS_DIR = ROOT / "reports"
IMAGE_DIR = ROOT / "01_Data" / "Example of Inventory" / "Images"
JSON_DIR = ROOT / "vllm_pipeline" / "vllm_output"
REPORT_PATH = REPORTS_DIR / "vllm_report.docx"
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp", ".heic"}
IMAGE_WIDTH_CM = 8.0


def image_stream(path: Path) -> io.BytesIO:
    img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    img.thumbnail((1200, 1200))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    buf.seek(0)
    return buf


def add_line(cell, label: str, value, first=False):
    if isinstance(value, (list, tuple)):
        value = ", ".join(str(v) for v in value)
    if not value:
        return
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    run = p.add_run(f"{label}: ")
    run.bold = True
    p.add_run(str(value))


def fill_analysis(cell, data: dict):
    summary = data.get("audit_summary", {})
    add_line(cell, "Space type", summary.get("space_type"), first=True)
    add_line(cell, "Overall condition", summary.get("overall_condition"))

    for comp in data.get("salvageable_components", []):
        p = cell.add_paragraph()
        run = p.add_run(f"{comp.get('component_id', '')} - {comp.get('item_name', '')}")
        run.bold = True
        run.font.size = Pt(10.5)
        tech = comp.get("technical_features", {})
        add_line(cell, "Category", comp.get("category"))
        add_line(cell, "Materials", comp.get("estimated_materials"))
        add_line(cell, "Salvage tier", comp.get("salvage_tier"))
        add_line(cell, "Hardware", tech.get("hardware_present"))
        add_line(cell, "Markings", tech.get("labels_or_markings"))
        add_line(cell, "Connection", tech.get("connection_type"))
        add_line(cell, "Deconstruction notes", comp.get("deconstruction_notes"))
        add_line(cell, "Bounding box [ymin, xmin, ymax, xmax] (0-1000)", comp.get("bounding_box_2d"))

    flags = data.get("hazard_and_diagnostic_flags", [])
    if flags:
        p = cell.add_paragraph()
        p.add_run("Hazard and diagnostic flags").bold = True
        for flag in flags:
            cell.add_paragraph(str(flag), style="List Bullet")


def main():
    images = sorted(
        (p for p in IMAGE_DIR.iterdir() if p.suffix.lower() in IMAGE_SUFFIXES),
        key=lambda p: p.name.lower(),
    )

    doc = Document()
    section = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))
    doc.styles["Normal"].font.size = Pt(9)
    doc.add_heading("Deconstruction audit - vLLM analysis", level=0)

    for path in images:
        doc.add_heading(path.name, level=1)
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        img_cell, text_cell = table.rows[0].cells
        img_cell.width, text_cell.width = Cm(IMAGE_WIDTH_CM + 0.5), Cm(9.5)

        img_cell.paragraphs[0].add_run().add_picture(image_stream(path), width=Cm(IMAGE_WIDTH_CM))

        json_path = JSON_DIR / f"{path.stem}.json"
        if not json_path.exists():
            text_cell.paragraphs[0].add_run("No vLLM output for this image.")
            continue
        try:
            fill_analysis(text_cell, json.loads(json_path.read_text(encoding="utf-8")))
        except json.JSONDecodeError:
            # keep the raw model answer so nothing is lost
            text_cell.paragraphs[0].add_run(json_path.read_text(encoding="utf-8"))

    doc.save(REPORT_PATH)
    print(f"Saved {REPORT_PATH}")


if __name__ == "__main__":
    main()
